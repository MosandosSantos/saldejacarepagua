"""Renderers in memory for the initiation-flow listing report (PDF/XLSX).

O relatório é a listagem dos irmãos selecionados pelos filtros do Flow de
Iniciação: cadastro, nome, grau, data do último grau e presença nas
chamadas. O cabeçalho identifica corpos, loja e vale selecionados.
"""

from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

_PDF_NAVY = colors.HexColor('#0a1e3d')
_PDF_TEXT = colors.HexColor('#243447')
_PDF_MUTED = colors.HexColor('#526477')
_PDF_GRID = colors.HexColor('#cbd5e1')
_XLSX_NAVY = '0A1E3D'
_XLSX_ACCENT = '5B8DEF'
_XLSX_GRID = 'CBD5E1'


def _texto_pdf(valor):
    texto = '' if valor is None else str(valor)
    texto = texto.encode('latin-1', 'replace').decode('latin-1')
    return escape(texto)


def _celula_texto(worksheet, row, column, value):
    cell = worksheet.cell(row=row, column=column, value='' if value is None else str(value))
    cell.data_type = 's'
    return cell


def _descritor(contexto):
    """Linha única com corpos, loja e vale do relatório."""
    partes = [
        f'Corpos: {corpo}' for corpo in contexto.get('corpos', [])
    ]
    if not partes:
        partes = ['Todos os corpos do Rito']
    if contexto.get('loja'):
        partes.append(f'Loja: {contexto["loja"]}')
    if contexto.get('vale'):
        partes.append(f'Vale: {contexto["vale"]}')
    return ' · '.join(partes)


def _formatar_data(valor):
    if valor is None:
        return '—'
    if isinstance(valor, datetime):
        valor = valor.date()
    return valor.strftime('%d/%m/%Y')


def _cabecalhos():
    return ['Cadastro', 'Nome do Irmão', 'Grau', 'Último grau', 'Presença']


def _linhas_tabela(linhas):
    registro = []
    for linha in linhas:
        cadastro = (
            str(linha['cadastro']) if linha['cadastro'] is not None else '—'
        )
        data_ultimo = _formatar_data(linha['data_ultimo'])
        presenca = f"{linha['presenca']}%"
        registro.append([
            cadastro,
            linha['nome'],
            f"{linha['grau']}º",
            data_ultimo,
            presenca,
        ])
    return registro


def gerar_xlsx_iniciacao(linhas, contexto, gerado_em):
    """Return an XLSX workbook in memory with the listing rows."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Iniciação'
    workbook.properties.title = 'Relatório de Iniciação'
    workbook.properties.subject = (
        f'{_descritor(contexto)} · gerado em {gerado_em:%d/%m/%Y}'
    )

    for column, value in enumerate(_cabecalhos(), start=1):
        cell = _celula_texto(worksheet, 1, column, value)
        cell.font = Font(name='Manrope', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=_XLSX_NAVY)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_number, registro in enumerate(_linhas_tabela(linhas), start=2):
        for column, value in enumerate(registro, start=1):
            _celula_texto(worksheet, row_number, column, value)

    ultima_coluna = get_column_letter(len(_cabecalhos()))
    worksheet.auto_filter.ref = f'A1:{ultima_coluna}{max(1, worksheet.max_row)}'
    worksheet.freeze_panes = 'C2'
    worksheet.sheet_view.showGridLines = True
    worksheet.sheet_view.zoomScale = 85
    worksheet.print_title_rows = '1:1'
    worksheet.page_setup.orientation = 'landscape'
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    for column, width in zip('ABCDE', (12, 40, 9, 13, 11), strict=True):
        worksheet.column_dimensions[column].width = width

    thin_grid = Side(style='thin', color=_XLSX_GRID)
    for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row):
        for cell in row:
            cell.border = Border(
                left=thin_grid, right=thin_grid, top=thin_grid, bottom=thin_grid,
            )
            if cell.row > 1:
                cell.alignment = Alignment(
                    horizontal=(
                        'left' if cell.column == 2 else 'center'
                    ),
                    vertical='center',
                )
                if cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F3F6FA')
                if cell.column == 5:
                    cell.font = Font(
                        color=(
                            _XLSX_ACCENT if not cell.value.startswith('-')
                            else 'DC2626'
                        ),
                        bold=True,
                    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def gerar_pdf_iniciacao(linhas, contexto, gerado_em):
    """Return a compact A4 (portrait) listing of the selected brothers."""
    page_width, _page_height = A4
    margem_horizontal = 8 * mm
    largura_util = page_width - 2 * margem_horizontal
    larguras = [0.13, 0.37, 0.07, 0.13, 0.10]
    larguras_colunas = [
        largura_util * (p / sum(larguras)) for p in larguras
    ]

    titulo_style = ParagraphStyle(
        'IniciacaoTitulo', fontName='Helvetica-Bold', fontSize=13,
        leading=15, textColor=_PDF_NAVY, spaceAfter=1.5 * mm,
    )
    subtitulo_style = ParagraphStyle(
        'IniciacaoSubtitulo', fontName='Helvetica', fontSize=7.5,
        leading=9, textColor=_PDF_MUTED, spaceAfter=4 * mm,
    )
    header_style = ParagraphStyle(
        'IniciacaoCabecalho', fontName='Helvetica-Bold', fontSize=7,
        leading=9, textColor=colors.white, alignment=1,
    )
    cell_style = ParagraphStyle(
        'IniciacaoValor', fontName='Helvetica', fontSize=7,
        leading=9, textColor=_PDF_TEXT, alignment=1,
    )
    nome_style = ParagraphStyle(
        'IniciacaoNome', fontName='Helvetica', fontSize=7,
        leading=9, textColor=_PDF_TEXT,
    )
    empty_style = ParagraphStyle(
        'IniciacaoVazio', fontName='Helvetica-Oblique', fontSize=7.5,
        leading=10, textColor=_PDF_MUTED, alignment=1,
    )

    dados = [[
        Paragraph(_texto_pdf(valor), header_style)
        for valor in _cabecalhos()
    ]]
    for registro in _linhas_tabela(linhas):
        dados.append([
            Paragraph(_texto_pdf(valor), nome_style if i == 1 else cell_style)
            for i, valor in enumerate(registro)
        ])

    estilo = [
        ('BACKGROUND', (0, 0), (-1, 0), _PDF_NAVY),
        ('GRID', (0, 0), (-1, -1), 0.25, _PDF_GRID),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('LEFTPADDING', (0, 0), (-1, -1), 2),
        ('RIGHTPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]
    if not linhas:
        dados.append([
            Paragraph('Nenhum irmão encontrado para esses filtros.', empty_style),
        ] + [''] * (len(_cabecalhos()) - 1))
        estilo.append(('SPAN', (0, 1), (-1, 1)))
    else:
        for row_number in range(2, len(dados), 2):
            estilo.append((
                'BACKGROUND', (0, row_number), (-1, row_number),
                colors.HexColor('#f3f6fa'),
            ))

    tabela = Table(
        dados,
        colWidths=larguras_colunas,
        repeatRows=1,
        splitByRow=1,
        hAlign='LEFT',
    )
    tabela.setStyle(TableStyle(estilo))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=margem_horizontal,
        rightMargin=margem_horizontal,
        topMargin=10 * mm,
        bottomMargin=11 * mm,
        title='Relatório de Iniciação',
        author='GPRBRCRM',
    )
    story = [
        Paragraph('Relatório de Iniciação', titulo_style),
        Paragraph(
            (
                f'{_descritor(contexto)} — {len(linhas)} irmão(s) — '
                f'gerado em {gerado_em:%d/%m/%Y}'
            ),
            subtitulo_style,
        ),
        tabela,
    ]
    doc.build(story)
    return buffer.getvalue()
