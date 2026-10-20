"""Pagination and PDF/XLSX export for the tenant-scoped Loja list."""
import re
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import parse_qs, urlparse
from uuid import uuid4

import pytest
from django.conf import settings
from django.contrib.auth.hashers import make_password
from django.test import Client
from openpyxl import load_workbook

from acesso.models import Tenant, Usuario
from acesso.session import iniciar_sessao
from contatos.models import Loja, Membro, Vale

pytestmark = pytest.mark.django_db


@pytest.fixture
def lojas_para_relatorio():
    sufixo = uuid4().hex
    tenant = Tenant.objects.create(
        nome='Vale Relatório Lojas', slug=f'lojas-relatorio-{sufixo}',
    )
    Vale.objects.create(
        tenant=tenant, nome='Vale Relatório Lojas', sigla=f'VR{sufixo[:4]}',
    )
    tenant_externo = Tenant.objects.create(
        nome='Vale Externo Lojas', slug=f'lojas-externas-{sufixo}',
    )
    Vale.objects.create(
        tenant=tenant_externo,
        nome='Vale Externo Lojas',
        sigla=f'VE{sufixo[:4]}',
    )
    presidente = Membro.objects.create(
        tenant=tenant,
        nome='Presidente das Lojas',
        celular='(21) 90000-0000',
    )

    lojas_filtradas = [
        Loja.objects.create(
            tenant=tenant,
            nome=f'Loja Relatorio {indice:02d}',
            corpo='perfeicao',
            deputado_mestre=presidente,
        )
        for indice in range(21)
    ]
    Loja.objects.create(
        tenant=tenant,
        nome='Loja Outro Corpo',
        corpo='kadosh',
    )
    Loja.objects.create(
        tenant=tenant_externo,
        nome='Loja Relatorio Sigilosa',
        corpo='perfeicao',
    )

    usuario = Usuario.objects.create(
        nome='Operador Relatório Lojas',
        email=f'lojas-{sufixo}@example.test',
        senha=make_password('senha123'),
        role='operacional',
        tenant=tenant,
    )
    cliente = Client()
    sessao = cliente.session
    iniciar_sessao(SimpleNamespace(session=sessao), usuario)
    sessao.save()
    cliente.cookies[settings.SESSION_COOKIE_NAME] = sessao.session_key
    return {
        'cliente': cliente,
        'tenant': tenant,
        'lojas': lojas_filtradas,
    }


def _filtros():
    return {'q': 'Loja Relatorio', 'corpo': ['perfeicao']}


def test_lista_lojas_pagina_em_vinte_preservando_filtros(lojas_para_relatorio):
    quadro = lojas_para_relatorio

    resposta = quadro['cliente'].get('/lojas/', _filtros())

    assert resposta.status_code == 200
    assert len(resposta.context['lojas']) == 20
    assert resposta.context['total'] == 21
    assert resposta.context['page_obj'].paginator.per_page == 20
    assert 'Página 1 de 2' in resposta.content.decode()
    assert 'Próxima' in resposta.content.decode()
    assert 'aria-label="Paginação"' in resposta.content.decode()
    assert any(
        page['number'] == 2 for page in resposta.context['pagination_urls']['pages']
    )
    url_pagina_2 = next(
        page['url'] for page in resposta.context['pagination_urls']['pages']
        if page['number'] == 2
    )
    parametros_pagina_2 = parse_qs(urlparse(url_pagina_2).query)
    assert parametros_pagina_2['q'] == ['Loja Relatorio']
    assert parametros_pagina_2['corpo'] == ['perfeicao']
    parametros_next = parse_qs(urlparse(
        resposta.context['pagination_urls']['next'],
    ).query)
    assert parametros_next['q'] == ['Loja Relatorio']
    assert parametros_next['corpo'] == ['perfeicao']

    segunda = quadro['cliente'].get(resposta.context['pagination_urls']['next'])
    assert segunda.status_code == 200
    assert len(segunda.context['lojas']) == 1
    assert segunda.context['lojas'][0]['nome'] == 'Loja Relatorio 20'
    assert 'Anterior' in segunda.content.decode()


def test_legacy_lojas_list_paginates_to_the_canonical_root(lojas_para_relatorio):
    response = lojas_para_relatorio['cliente'].get('/contatos/lojas/', _filtros())

    assert response.status_code == 200
    assert response.context['pagination_urls']['next'].startswith('/lojas/?')


def test_lista_lojas_corpo_invalido_nao_amplia_resultados(lojas_para_relatorio):
    quadro = lojas_para_relatorio

    resposta = quadro['cliente'].get('/contatos/lojas/', {
        'q': 'Loja Relatorio',
        'corpo': ['corpo-invalido'],
    })

    assert resposta.status_code == 200
    assert resposta.context['total'] == 0
    assert resposta.context['lojas'] == []


def test_xlsx_exporta_todas_as_lojas_filtradas_sem_vazar_tenant(lojas_para_relatorio):
    quadro = lojas_para_relatorio

    resposta = quadro['cliente'].get(
        '/lojas/relatorio/xlsx/',
        {**_filtros(), 'page': '2'},
    )

    assert resposta.status_code == 200
    assert resposta['Content-Type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    assert resposta['Content-Disposition'].endswith('.xlsx"')
    planilha = load_workbook(BytesIO(resposta.content), data_only=False)['Lojas']
    assert planilha.max_row == 22
    assert planilha.auto_filter.ref == 'A1:G22'
    assert planilha.freeze_panes == 'A2'
    assert [celula.value for celula in planilha[1]] == [
        'Loja', 'Corpo', 'Vale', 'Presidente', 'Celular do presidente',
        'Membros', 'Situação',
    ]
    nomes = {planilha.cell(row=row, column=1).value for row in range(2, 23)}
    assert nomes == {loja.nome for loja in quadro['lojas']}
    assert 'Loja Outro Corpo' not in nomes
    assert 'Loja Relatorio Sigilosa' not in nomes


def test_pdf_exporta_relatorio_a4_paisagem(lojas_para_relatorio):
    quadro = lojas_para_relatorio
    Loja.objects.bulk_create([
        Loja(
            tenant=quadro['tenant'],
            nome=f'Loja Relatorio Extra {indice:03d}',
            corpo='perfeicao',
        )
        for indice in range(100)
    ])

    resposta = quadro['cliente'].get(
        '/contatos/lojas/relatorio/pdf/',
        {**_filtros(), 'page': '2'},
    )

    assert resposta.status_code == 200
    assert resposta['Content-Type'] == 'application/pdf'
    assert resposta['Content-Disposition'].endswith('.pdf"')
    assert resposta.content.startswith(b'%PDF-')
    media_box = re.search(
        rb'/MediaBox\s*\[\s*0\s+0\s+([\d.]+)\s+([\d.]+)\s*\]',
        resposta.content,
    )
    assert media_box is not None
    assert float(media_box.group(1)) > float(media_box.group(2))
    assert len(re.findall(rb'/Type\s*/Page\b', resposta.content)) > 1


def test_relatorio_rejeita_formato_desconhecido(lojas_para_relatorio):
    quadro = lojas_para_relatorio

    resposta = quadro['cliente'].get('/contatos/lojas/relatorio/csv/')

    assert resposta.status_code == 400
    assert 'Content-Disposition' not in resposta


def test_toolbar_mostra_botao_relatorio_com_pdf_e_excel(lojas_para_relatorio):
    quadro = lojas_para_relatorio

    resposta = quadro['cliente'].get('/lojas/', _filtros())
    html = resposta.content.decode()

    assert resposta.status_code == 200
    assert 'Relatório' in html
    assert 'PDF' in html
    assert 'Excel (.xlsx)' in html
    assert 'formaction="/lojas/relatorio/pdf/"' in html
    assert 'formaction="/lojas/relatorio/xlsx/"' in html
    assert 'name="q" value="Loja Relatorio"' in html
    assert 'name="corpo" value="perfeicao" checked' in html
    assert 'name="page"' not in html


def test_exportacao_vazia_ainda_entrega_arquivos_com_cabecalho():
    sufixo = uuid4().hex[:10]
    tenant = Tenant.objects.create(
        nome='Vale Lojas Vazio', slug=f'lojas-vazias-{sufixo}',
    )
    Vale.objects.create(
        tenant=tenant, nome='Vale Lojas Vazio', sigla=f'VV{sufixo[:4]}',
    )
    usuario = Usuario.objects.create(
        nome='Operador Lojas Vazio',
        email=f'lojas-vazias-{uuid4().hex}@example.test',
        senha=make_password('senha123'),
        role='operacional',
        tenant=tenant,
    )
    cliente = Client()
    sessao = cliente.session
    iniciar_sessao(SimpleNamespace(session=sessao), usuario)
    sessao.save()
    cliente.cookies[settings.SESSION_COOKIE_NAME] = sessao.session_key

    resposta = cliente.get('/contatos/lojas/relatorio/xlsx/', {'q': 'sem resultados'})

    assert resposta.status_code == 200
    planilha = load_workbook(BytesIO(resposta.content), data_only=False)['Lojas']
    assert planilha.max_row == 1
    assert [celula.value for celula in planilha[1]][:2] == ['Loja', 'Corpo']

    pdf_vazio = cliente.get('/contatos/lojas/relatorio/pdf/', {'q': 'sem resultados'})
    assert pdf_vazio.status_code == 200
    assert pdf_vazio.content.startswith(b'%PDF-')


def test_xlsx_mantem_nome_com_prefixo_de_formula_como_texto(lojas_para_relatorio):
    quadro = lojas_para_relatorio
    Loja.objects.create(
        tenant=quadro['tenant'], nome='=2+2', corpo='perfeicao',
    )

    resposta = quadro['cliente'].get(
        '/contatos/lojas/relatorio/xlsx/', {'q': '=2+2'},
    )

    assert resposta.status_code == 200
    planilha = load_workbook(BytesIO(resposta.content), data_only=False)['Lojas']
    celula = planilha['A2']
    assert celula.value == '=2+2'
    assert celula.data_type == 's'


def test_exportacao_bloqueia_usuario_sem_vale_ativo():
    sufixo = uuid4().hex
    tenant = Tenant.objects.create(
        nome='Tenant sem Vale', slug=f'lojas-sem-vale-{sufixo}',
    )
    usuario = Usuario.objects.create(
        nome='Operador sem Vale',
        email=f'lojas-sem-vale-{sufixo}@example.test',
        senha=make_password('senha123'),
        role='operacional',
        tenant=tenant,
    )
    cliente = Client()
    sessao = cliente.session
    iniciar_sessao(SimpleNamespace(session=sessao), usuario)
    sessao.save()
    cliente.cookies[settings.SESSION_COOKIE_NAME] = sessao.session_key

    resposta = cliente.get('/contatos/lojas/relatorio/xlsx/')

    assert resposta.status_code == 403
    assert 'Content-Disposition' not in resposta
