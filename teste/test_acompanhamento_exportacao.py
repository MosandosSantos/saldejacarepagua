import re
from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from django.contrib.auth.hashers import make_password
from django.core.management import call_command
from django.test import Client
from openpyxl import load_workbook

from acesso.models import ROLE_DELEGADO, Tenant, Usuario
from contatos.models import Grau, GrauMembro, Loja, Membro, Vale


def _tenant_vale(nome, sigla):
    suffix = uuid4().hex
    tenant = Tenant.objects.create(nome=nome, slug=f'{sigla.lower()}-{suffix}')
    vale = Vale.objects.create(
        tenant=tenant, nome=nome, sigla=f'{sigla}{suffix[:4]}',
    )
    return tenant, vale


def _cliente(tenant):
    suffix = uuid4().hex
    usuario = Usuario.objects.create(
        nome='Usuário de exportação',
        email=f'exportacao-{suffix}@example.test',
        senha=make_password('senha-teste'),
        role=ROLE_DELEGADO,
        tenant=tenant,
    )
    client = Client()
    session = client.session
    session['usuario_id'] = usuario.pk
    session['usuario'] = {
        'nome': usuario.nome,
        'email': usuario.email,
        'role': usuario.role,
        'tenant_id': tenant.pk,
    }
    session.save()
    return client


def _registrar_grau(membro, numero, data):
    grau, _ = Grau.objects.get_or_create(
        numero=numero, defaults={'nome': f'Grau {numero}'},
    )
    return GrauMembro.objects.create(
        tenant_id=membro.tenant_id,
        membro=membro,
        grau=grau,
        data=data,
    )


@pytest.mark.django_db
def test_xlsx_exporta_todos_os_filtrados_sem_vazar_tenant_ou_limitar_pagina():
    call_command('seed_graus', verbosity=0)
    tenant_a, vale_a = _tenant_vale('Vale Acompanhamento A', 'VXA')
    tenant_b, _ = _tenant_vale('Vale Acompanhamento B', 'VXB')
    loja_a = Loja.objects.create(
        tenant=tenant_a, nome='Loja de Perfeição A', corpo='perfeicao',
    )
    loja_b = Loja.objects.create(
        tenant=tenant_b, nome='Loja de Perfeição B', corpo='perfeicao',
    )

    nomes_a = []
    for indice in range(21):
        membro = Membro.objects.create(
            tenant=tenant_a, nome=f'Membro Exportado {indice:02d}',
        )
        membro.lojas.add(loja_a)
        _registrar_grau(membro, 10, date(2010, 3, 14))
        nomes_a.append(membro.nome)

    membro_b = Membro.objects.create(tenant=tenant_b, nome='Membro Sigiloso B')
    membro_b.lojas.add(loja_b)
    _registrar_grau(membro_b, 10, date(2010, 3, 14))
    client = _cliente(tenant_a)

    response = client.get('/membros/acompanhamento/relatorio/xlsx/', {
        'vale': str(vale_a.pk),
        'loja': str(loja_a.pk),
        'grupo_grau': 'perfeicao',
        'page': '2',
    })

    assert response.status_code == 200
    assert response['Content-Type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    assert response['Content-Disposition'].endswith('.xlsx"')

    workbook = load_workbook(BytesIO(response.content), data_only=False)
    sheet = workbook['Acompanhamento']
    assert sheet.max_row == 22
    assert sheet.freeze_panes == 'C2'
    assert sheet.auto_filter.ref == 'A1:M22'
    cabecalhos = [cell.value for cell in sheet[1]]
    assert cabecalhos[:2] == ['Membro', 'Presença']
    assert len(cabecalhos) == 13  # Membro, Presença e graus 4º–14º
    indice_grau_10 = cabecalhos.index('10º') + 1
    data_grau = sheet.cell(row=2, column=indice_grau_10)
    assert (data_grau.value.year, data_grau.value.month, data_grau.value.day) == (2010, 3, 14)
    assert data_grau.number_format == 'mm/yyyy'
    nomes_exportados = {sheet.cell(row=row, column=1).value for row in range(2, 23)}
    assert nomes_exportados == set(nomes_a)
    assert 'Membro Sigiloso B' not in nomes_exportados


@pytest.mark.django_db
def test_exportacao_bloqueia_tenant_sem_vale_ativo():
    tenant = Tenant.objects.create(
        nome='Tenant sem Vale ativo', slug=f'sem-vale-{uuid4().hex}',
    )
    response = _cliente(tenant).get(
        '/membros/acompanhamento/relatorio/xlsx/',
    )

    assert response.status_code == 403
    assert 'Content-Disposition' not in response


@pytest.mark.django_db
def test_exportacao_nao_aceita_formato_desconhecido():
    tenant, _ = _tenant_vale('Vale Formato Inválido', 'VFI')

    response = _cliente(tenant).get(
        '/membros/acompanhamento/relatorio/csv/',
    )

    assert response.status_code == 400
    assert 'Content-Disposition' not in response


@pytest.mark.django_db
def test_exportacao_nao_aceita_vale_de_outro_tenant():
    tenant_a, _ = _tenant_vale('Vale Solicitante', 'VSO')
    tenant_b, vale_b = _tenant_vale('Vale Não Autorizado', 'VNA')
    membro_b = Membro.objects.create(tenant=tenant_b, nome='Sigiloso do Vale B')
    _registrar_grau(membro_b, 10, date(2010, 3, 14))

    response = _cliente(tenant_a).get(
        '/membros/acompanhamento/relatorio/xlsx/',
        {'vale': str(vale_b.pk)},
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    sheet = workbook['Acompanhamento']
    assert sheet.max_row == 1
    assert all(cell.value != 'Sigiloso do Vale B' for row in sheet.iter_rows() for cell in row)


@pytest.mark.django_db
def test_xlsx_mantem_nome_com_prefixo_de_formula_como_texto():
    call_command('seed_graus', verbosity=0)
    tenant, _ = _tenant_vale('Vale Fórmula', 'VFO')
    membro = Membro.objects.create(tenant=tenant, nome='=2+2')
    _registrar_grau(membro, 10, date(2010, 3, 14))

    response = _cliente(tenant).get('/membros/acompanhamento/relatorio/xlsx/')

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    cell = workbook['Acompanhamento']['A2']
    assert cell.value == '=2+2'
    assert cell.data_type == 's'


@pytest.mark.django_db
def test_xlsx_sem_resultados_ainda_entrega_cabecalho():
    tenant, _ = _tenant_vale('Vale Vazio', 'VVA')

    response = _cliente(tenant).get('/membros/acompanhamento/relatorio/xlsx/')

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    sheet = workbook['Acompanhamento']
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1][:2]] == ['Membro', 'Presença']


@pytest.mark.django_db
def test_exportacao_ignora_historico_de_grau_de_outro_tenant():
    call_command('seed_graus', verbosity=0)
    tenant_a, _ = _tenant_vale('Vale do Membro', 'VMM')
    tenant_b, _ = _tenant_vale('Vale do Histórico', 'VMH')
    membro = Membro.objects.create(tenant=tenant_a, nome='Membro local')
    grau_33 = Grau.objects.get(numero=33)
    GrauMembro.objects.create(
        tenant=tenant_b, membro=membro, grau=grau_33, data=date(2010, 3, 14),
    )
    client = _cliente(tenant_a)

    completo = client.get('/membros/acompanhamento/relatorio/xlsx/')
    filtrado = client.get(
        '/membros/acompanhamento/relatorio/xlsx/', {'grupo_grau': 'sc33'},
    )

    assert completo.status_code == filtrado.status_code == 200
    workbook_completo = load_workbook(BytesIO(completo.content), data_only=False)
    sheet_completo = workbook_completo['Acompanhamento']
    grau_33_col = [cell.value for cell in sheet_completo[1]].index('33º') + 1
    assert sheet_completo.cell(row=2, column=grau_33_col).value == '—'
    workbook_filtrado = load_workbook(BytesIO(filtrado.content), data_only=False)
    assert workbook_filtrado['Acompanhamento'].max_row == 1


@pytest.mark.django_db
def test_pdf_exporta_tabela_em_a4_paisagem():
    call_command('seed_graus', verbosity=0)
    tenant, vale = _tenant_vale('Vale PDF Paisagem', 'VPP')
    membro = Membro.objects.create(tenant=tenant, nome='Membro para PDF')
    _registrar_grau(membro, 10, date(2010, 3, 14))
    response = _cliente(tenant).get(
        '/membros/acompanhamento/relatorio/pdf/', {'vale': str(vale.pk)},
    )

    assert response.status_code == 200
    assert response['Content-Type'] == 'application/pdf'
    assert response['Content-Disposition'].endswith('.pdf"')
    assert response.content.startswith(b'%PDF-')
    media_box = re.search(
        rb'/MediaBox\s*\[\s*0\s+0\s+([\d.]+)\s+([\d.]+)\s*\]',
        response.content,
    )
    assert media_box is not None
    assert float(media_box.group(1)) > float(media_box.group(2))


@pytest.mark.django_db
def test_toolbar_tem_icones_e_envia_filtros_atuais_para_relatorio():
    tenant, vale = _tenant_vale('Vale Toolbar Relatório', 'VTR')
    loja = Loja.objects.create(tenant=tenant, nome='Loja Toolbar')
    response = _cliente(tenant).get('/membros/acompanhamento/', {
        'vale': str(vale.pk),
        'loja': str(loja.pk),
        'grupo_grau': 'kadosh',
        'grau': ['19'],
        'ordenar': 'grau',
        'direcao': 'desc',
        'page': '2',
    })
    html = response.content.decode()
    form_start = html.index('<form method="get"')
    form_end = html.index('</form>', form_start) + len('</form>')
    form = html[form_start:form_end]

    assert response.status_code == 200
    assert 'Filtrar' in form
    assert 'Limpar' in form
    assert 'Relatório' in form
    assert form.count('<svg') >= 3
    assert form.count('aria-hidden="true"') >= 3
    assert 'formaction="/membros/acompanhamento/relatorio/pdf/"' in form
    assert 'formaction="/membros/acompanhamento/relatorio/xlsx/"' in form
    assert 'name="vale"' in form and f'value="{vale.pk}"' in form
    assert 'name="loja"' in form and f'value="{loja.pk}"' in form
    assert 'name="grupo_grau" value="kadosh"' in form
    assert 'name="grau" value="19"' in form
    assert 'name="ordenar" value="grau"' in form
    assert 'name="direcao" value="desc"' in form
    assert 'name="page"' not in form
