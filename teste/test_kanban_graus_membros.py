"""Colunas de grau do Flow de Iniciação — cadeia de marcos com listagem
de TODOS os detentores do grau anterior, com tempo mínimo de permanência.

Regras (marco → quem entra → tempo mínimo do grau âncora):
- 'Grau 9' (Iniciação G9): todos do grau 8 com 6+ meses do grau 4;
- 'Grau 14': todos do grau 13 com 3+ meses do grau 9;
- 'Grau 15': todos do grau 14 com 6+ meses do grau 14;
- 'Grau 18' (ex-'Comunicação'): todos do grau 17 com 6+ meses do grau 15;
- 'Grau 19': todos do grau 18 com 6+ meses do grau 18;
- 'Grau 22' (ex-'Diploma'): todos do grau 21 com 4+ meses do grau 19;
- 'Grau 28': todos do grau 27 com 4+ meses do grau 22;
- 'Grau 30': todos do grau 29 com 4+ meses do grau 28;
- 'Grau 31': todos do grau 30 com 6+ meses do grau 30;
- 'Grau 32': todos do grau 31 com 3+ meses do grau 31.

Cadência: 18→19 (6m), 19→22 (4m), 22→28 (4m), 28→30 (4m),
30→31 (6m), 31→32 (3m).
Todos inclui ativos e inativos; sem âncora ou tempo insuficiente, fora.

Card de membro: sem o nome do Vale, com o percentual de presença nas
chamadas — nome em vermelho quando abaixo de 50%.
"""
from datetime import timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from django.contrib.auth.hashers import make_password
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from acesso.models import Tenant, Usuario
from contatos.models import Grau, GrauMembro, Loja, Membro, Vale
from tarefas.models import Evento, PresencaEvento
from tarefas.presencas import PresencaServico

pytestmark = pytest.mark.django_db(transaction=True)


def _cliente(usuario):
    client = Client()
    session = client.session
    session['usuario_id'] = usuario.id
    session['usuario'] = {
        'nome': usuario.nome,
        'email': usuario.email,
        'role': usuario.role,
        'tenant_id': usuario.tenant_id,
    }
    session['nicho'] = 'saude'
    session.save()
    return client


def _admin():
    return Usuario.objects.create(
        nome='Admin Graus',
        email=f'admin-graus-{uuid4().hex}@example.test',
        role='admin_saas',
        tenant=None,
        senha=make_password('test-password'),
    )


def _mundo():
    sufixo = uuid4().hex
    tenant = Tenant.objects.create(
        nome='Vale Graus Teste', slug=f'vale-graus-{sufixo}',
    )
    Vale.objects.create(
        tenant=tenant, nome=f'Vale Aurora {sufixo}',
        sigla=f'VA{sufixo[:6].upper()}',
    )
    loja = Loja.objects.create(
        tenant=tenant, nome=f'Loja Perfeição {sufixo}', corpo='perfeicao',
    )
    graus = {}
    for n in (4, 8, 9, 13, 14, 15, 17, 18, 19, 21, 22, 27, 28, 29, 30, 31):
        graus[n] = Grau.objects.get_or_create(
            numero=n, defaults={'nome': f'Grau {n}'},
        )[0]

    hoje = timezone.localdate()

    def meses(n):
        return hoje - timedelta(days=30 * n)

    def membro(nome, graus_datas, situacao='ATIVO'):
        """graus_datas: lista de (nº do grau, data de concessão)."""
        m = Membro.objects.create(
            tenant=tenant, nome=nome, situacao=situacao,
        )
        for numero, data in graus_datas:
            GrauMembro.objects.create(
                tenant=tenant, membro=m, grau=graus[numero], data=data,
            )
        m.lojas.add(loja)
        return m

    return {
        'tenant': tenant,
        'loja': loja,
        'hoje': hoje,
        # G9 (Iniciação G9): todos do grau 8, 6+ meses do grau 4
        'g9_ok': membro(
            'Miro Grau Oito', [(4, meses(7)), (8, meses(4))],
        ),
        'g9_ok_inativo': membro(
            'Eva Grau Oito', [(4, meses(7)), (8, meses(4))],
            situacao='INATIVO',
        ),
        'g9_tempo_curto': membro(
            'Bruno Grau Oito', [(4, meses(5)), (8, meses(2))],
        ),
        'g9_sem_ancora': membro('Saulo Sem Suplica', [(8, meses(4))]),
        # G14: todos do grau 13, 3+ meses do grau 9
        'g14_ok': membro(
            'Ana Grau Treze', [(9, meses(4)), (13, meses(2))],
        ),
        'g14_tempo_curto': membro(
            'Pedro Grau Treze', [(9, meses(2)), (13, meses(1))],
        ),
        # G15: todos do grau 14, 6+ meses do grau 14
        'g15_ok': membro('Boa Grau Catorze', [(14, meses(7))]),
        # G18 (ex-Comunicação): todos do grau 17, 6+ meses do grau 15
        'g18_ok': membro(
            'Clara Grau Dezessete', [(15, meses(7)), (17, meses(3))],
        ),
        # G19: todos do grau 18, 6+ meses do grau 18
        'g19_ok': membro('Nuno Grau Dezoito', [(18, meses(7))]),
        'g19_tempo_curto': membro('Otto Grau Dezoito', [(18, meses(5))]),
        # G22 (ex-Diploma): todos do grau 21, 4+ meses do grau 19
        'g22_ok': membro(
            'Diego Grau Vinte e Um',
            [(18, meses(7)), (19, meses(5)), (21, meses(2))],
        ),
        'g22_tempo_curto': membro(
            'Paulo Grau Vinte e Um', [(19, meses(3)), (21, meses(1))],
        ),
        # G28: todos do grau 27, 4+ meses do grau 22
        'g28_ok': membro('Rita Grau Vinte e Sete', [
            (22, meses(5)), (27, meses(2)),
        ]),
        # G30: todos do grau 29, 4+ meses do grau 28
        'g30_ok': membro('Téo Grau Vinte e Nove', [
            (28, meses(5)), (29, meses(2)),
        ]),
        # G31: todos do grau 30, 6+ meses do grau 30
        'g31_ok': membro('Ulisses Grau Trinta', [(30, meses(7))]),
        'g31_tempo_curto': membro('Walter Grau Trinta', [(30, meses(5))]),
        # G32: todos do grau 31, 3+ meses do grau 31
        'g32_ok': membro('Xisto Grau Trinta e Um', [(31, meses(4))]),
        'g32_tempo_curto': membro('Yuri Grau Trinta e Um', [(31, meses(2))]),
    }


def _presencas(membro, loja, resultados):
    """Chamadas registradas: resultados = lista de bool (True = presente)."""
    for i, presente in enumerate(resultados):
        evento = Evento.objects.create(
            titulo=f'Sessão {i} — {membro.nome}',
            data_inicio=timezone.now(),
            data_fim=timezone.now(),
            loja_id=str(loja.pk),
        )
        PresencaEvento.objects.create(
            evento=evento, membro=membro, presente=presente,
        )


def test_marco_de_grau_respeita_grau_anterior_e_tempo_minimo():
    mundo = _mundo()
    client = _cliente(_admin())

    resposta = client.get('/iniciacao/')

    assert resposta.status_code == 200
    colunas = {c['key']: c for c in resposta.context['colunas']}

    # Cadeia completa de marcos (ex-Diploma agora é 'Grau 22')
    assert list(colunas) == [
        'suplica', 'pagamento', 'iniciacao', 'grau15', 'comunicacao',
        'grau19', 'diploma', 'grau28', 'grau30', 'grau31', 'grau32',
    ]
    assert colunas['comunicacao']['name'] == 'Grau 18'
    assert colunas['grau19']['name'] == 'Grau 19'
    assert colunas['diploma']['name'] == 'Grau 22'
    assert colunas['grau28']['name'] == 'Grau 28'
    assert colunas['grau30']['name'] == 'Grau 30'
    assert colunas['grau31']['name'] == 'Grau 31'
    assert colunas['grau32']['name'] == 'Grau 32'
    # Consistório (31º-32º) é o corpo das duas novas colunas
    assert colunas['grau31']['corpo'] == 'consistorio'
    assert colunas['grau32']['corpo'] == 'consistorio'

    esperado = {
        'pagamento': [mundo['g9_ok'], mundo['g9_ok_inativo']],
        'iniciacao': [mundo['g14_ok']],
        'grau15': [mundo['g15_ok']],
        'comunicacao': [mundo['g18_ok']],
        'grau19': [mundo['g19_ok']],
        'diploma': [mundo['g22_ok']],
        'grau28': [mundo['g28_ok']],
        'grau30': [mundo['g30_ok']],
        'grau31': [mundo['g31_ok']],
        'grau32': [mundo['g32_ok']],
    }
    for etapa, membros in esperado.items():
        cards_membros = [
            c for c in colunas[etapa]['cards'] if c.get('tipo') == 'membro'
        ]
        ids_esperados = [m.pk for m in sorted(membros, key=lambda m: m.nome)]
        assert [c['id'] for c in cards_membros] == ids_esperados

    # Inativo entra na listagem (todos os membros) e carrega a marcação
    cards_g9 = [
        c for c in colunas['pagamento']['cards'] if c.get('tipo') == 'membro'
    ]
    por_nome = {c['nome']: c for c in cards_g9}
    assert por_nome['Eva Grau Oito']['situacao'] == 'INATIVO'
    assert por_nome['Miro Grau Oito']['situacao'] == 'ATIVO'

    # Tempo insuficiente / sem âncora: fora da G9
    nomes_g9 = [c['nome'] for c in cards_g9]
    assert 'Bruno Grau Oito' not in nomes_g9
    assert 'Saulo Sem Suplica' not in nomes_g9

    # Tempo insuficiente nas demais listagens: fora
    cards_outros = [
        c for etapa in ('iniciacao', 'grau15', 'comunicacao', 'grau19',
                        'diploma', 'grau28', 'grau30', 'grau31', 'grau32')
        for c in colunas[etapa]['cards'] if c.get('tipo') == 'membro'
    ]
    nomes_outros = [c['nome'] for c in cards_outros]
    for nome in ('Pedro Grau Treze', 'Otto Grau Dezoito',
                 'Paulo Grau Vinte e Um', 'Walter Grau Trinta',
                 'Yuri Grau Trinta e Um'):
        assert nome not in nomes_outros


def test_filtro_vale_e_loja_no_quadro():
    """Vale filtra pelos membros do tenant do Vale; Loja filtra pelos
    membros vinculados à Loja (só Lojas daquele Vale quando um Vale está
    escolhido)."""
    mundo = _mundo()

    # Segundo vale + lojas próprias e cruzadas
    sufixo2 = uuid4().hex
    tenant_b = Tenant.objects.create(
        nome='Vale Graus B', slug=f'vale-graus-b-{sufixo2}',
    )
    Vale.objects.create(
        tenant=tenant_b, nome=f'Vale Borealis {sufixo2}',
        sigla=f'VB{sufixo2[:6].upper()}',
    )
    loja_b = Loja.objects.create(
        tenant=tenant_b, nome=f'Loja Harmonia {sufixo2}', corpo='perfeicao',
    )
    membro_b = Membro.objects.create(
        tenant=tenant_b, nome='Zeno Grau Oito Vale B', situacao='ATIVO',
    )
    for numero in (4, 8):
        GrauMembro.objects.create(
            tenant=tenant_b, membro=membro_b,
            grau=Grau.objects.get_or_create(
                numero=numero, defaults={'nome': f'Grau {numero}'},
            )[0],
            data=mundo['hoje'] - timedelta(days=30 * 7),
        )
    membro_b.lojas.add(loja_b)
    # Loja extra no Vale A (garante que o select do Vale B não a liste)
    Loja.objects.create(
        tenant=mundo['tenant'], nome=f'Loja Fenix {sufixo2}',
        corpo='capitulo',
    )

    client = _cliente(_admin())

    # 1) Sem filtro: membros dos dois vales aparecem
    colunas = {
        c['key']: c
        for c in client.get('/iniciacao/').context['colunas']
    }
    ids_g9 = [
        c['id'] for c in colunas['pagamento']['cards']
        if c.get('tipo') == 'membro'
    ]
    assert mundo['g9_ok'].pk in ids_g9
    assert membro_b.pk in ids_g9

    # 2) Vale B: só membros do tenant do Vale B
    resposta_b = client.get('/iniciacao/', {
        'prefeitura': Vale.objects.get(tenant=tenant_b).pk,
    })
    colunas_b = {c['key']: c for c in resposta_b.context['colunas']}
    ids_b = [
        c['id'] for c in colunas_b['pagamento']['cards']
        if c.get('tipo') == 'membro'
    ]
    assert ids_b == [membro_b.pk]
    # O select de Loja agora lista apenas as Lojas do Vale B (agrupadas por corpo)
    nomes_lojas_b = [
        loja['nome']
        for grupo in resposta_b.context['lojas_filtro']
        for loja in grupo['lojas']
    ]
    assert loja_b.nome in nomes_lojas_b
    assert not any('Loja Fenix' in nome for nome in nomes_lojas_b)

    # 3) Loja do Vale A: só membros vinculados àquela loja
    resposta_loja = client.get('/iniciacao/', {
        'loja': mundo['loja'].nome,
    })
    colunas_loja = {c['key']: c for c in resposta_loja.context['colunas']}
    ids_loja = [
        c['id'] for c in colunas_loja['pagamento']['cards']
        if c.get('tipo') == 'membro'
    ]
    assert mundo['g9_ok'].pk in ids_loja
    assert membro_b.pk not in ids_loja

    # 4) Loja inexistente: quadro sem cards de membro
    resposta_nada = client.get('/iniciacao/', {'loja': 'Loja Fantasma 999'})
    colunas_nada = {c['key']: c for c in resposta_nada.context['colunas']}
    assert not any(
        c.get('tipo') == 'membro'
        for col in colunas_nada.values()
        for c in col['cards']
    )


def test_filtro_esconde_colunas_sem_cards():
    """Com filtro de Vale/corpo ativo, coluna sem nenhum card sai do quadro;
    com filtro por Loja, a faixa do corpo fica fixa (coluna vazia permanece)."""
    mundo = _mundo()
    client = _cliente(_admin())

    # Uma loja de Kadosh no Vale A garante o grupo no select de Loja
    Loja.objects.create(
        tenant=mundo['tenant'], nome='Loja Kadosh Teste', corpo='kadosh',
    )

    # Loja é de Perfeição (4º-14º): a faixa do corpo fica FIXA — Grau 4/9/14
    # sempre, mesmo sem card (e nada do 15º adiante em loja de Perfeição)
    resposta_loja = client.get('/iniciacao/', {'loja': mundo['loja'].nome})
    chaves_loja = [c['key'] for c in resposta_loja.context['colunas']]
    assert 'grau15' not in chaves_loja
    assert 'comunicacao' not in chaves_loja
    assert 'grau31' not in chaves_loja
    assert chaves_loja == ['suplica', 'pagamento', 'iniciacao']

    # Capítulo: Grau 15 e 18 fixos, mesmo com a coluna do 15º vazia
    loja_capitulo = Loja.objects.create(
        tenant=mundo['tenant'], nome='Loja Capítulo Teste', corpo='capitulo',
    )
    resposta_cap = client.get(
        '/iniciacao/', {'loja': loja_capitulo.nome},
    )
    assert [c['key'] for c in resposta_cap.context['colunas']] == [
        'grau15', 'comunicacao',
    ]

    # Vale B com um único membro (grau 8) → só a coluna dele fica no quadro
    sufixo = uuid4().hex
    tenant_b = Tenant.objects.create(
        nome='Vale Graus B', slug=f'vale-graus-b-{sufixo}',
    )
    vale_b = Vale.objects.create(
        tenant=tenant_b, nome=f'Vale Borealis {sufixo}',
        sigla=f'VB{sufixo[:6].upper()}',
    )
    membro_b = Membro.objects.create(
        tenant=tenant_b, nome='Zeno Grau Oito Vale B', situacao='ATIVO',
    )
    for numero in (4, 8):
        GrauMembro.objects.create(
            tenant=tenant_b, membro=membro_b,
            grau=Grau.objects.get_or_create(
                numero=numero, defaults={'nome': f'Grau {numero}'},
            )[0],
            data=mundo['hoje'] - timedelta(days=30 * 7),
        )
    resposta_b = client.get('/iniciacao/', {'prefeitura': vale_b.pk})
    assert [c['key'] for c in resposta_b.context['colunas']] == ['pagamento']

    # Sem filtro: todas as colunas de grau presentes
    todas = [c['key'] for c in client.get('/iniciacao/').context['colunas']]
    assert 'suplica' in todas
    assert 'grau30' in todas

    # Faixas dos Corpos do Rito: Kadosh cobre as colunas 19/22/28/30
    kadosh = next(
        c for c in resposta_loja.context['corpos_rito']
        if c['valor'] == 'kadosh'
    )
    assert kadosh['faixa'] == '19º · 22º · 28º · 30º'
    html = resposta_loja.content.decode()
    # Select de Loja agrupado por corpo; botões Filtrar e Limpar não existem
    assert '<optgroup label="Kadosh · 19º · 22º · 28º · 30º">' in html
    assert 'Filtrar' not in html
    assert 'Limpar' not in html
    assert 'Filtrar' not in html
    assert 'Limpar' not in html


def test_filtro_corpo_lista_marcos_de_todos_os_vales():
    """Chips de Corpos do Rito: o quadro mostra os membros que serão
    iniciados nos graus do corpo — de TODOS os Vales e Lojas."""
    mundo = _mundo()
    client = _cliente(_admin())

    # Segundo Vale com membro do Capítulo (grau 17 → marco 18)
    sufixo = uuid4().hex
    tenant_b = Tenant.objects.create(
        nome='Vale Corpos B', slug=f'vale-corpos-b-{sufixo}',
    )
    Vale.objects.create(
        tenant=tenant_b, nome=f'Vale Corpos {sufixo}',
        sigla=f'VC{sufixo[:6].upper()}',
    )
    Loja.objects.create(
        tenant=tenant_b, nome=f'Loja de Capítulo corpos {sufixo}',
        corpo='capitulo',
    )
    zeno = Membro.objects.create(
        tenant=tenant_b, nome='Zeno Corpos Capitulo', situacao='ATIVO',
    )
    for numero in (15, 17):
        GrauMembro.objects.create(
            tenant=tenant_b, membro=zeno,
            grau=Grau.objects.get_or_create(
                numero=numero, defaults={'nome': f'Grau {numero}'},
            )[0],
            data=mundo['hoje'] - timedelta(days=30 * 7),
        )

    # Capítulo (15º-18º): colunas 15 e 18 com membros dos DOIS vales
    resposta = client.get('/iniciacao/', {'corpo': ['capitulo']})
    assert resposta.status_code == 200
    colunas = {c['key']: c for c in resposta.context['colunas']}
    assert list(colunas) == ['grau15', 'comunicacao']
    assert resposta.context['corpos_escolhidos'] == ['capitulo']
    ids_18 = [
        c['id'] for c in colunas['comunicacao']['cards']
        if c.get('tipo') == 'membro'
    ]
    assert ids_18 == sorted([
        mundo['g18_ok'].pk, zeno.pk,
    ], key=lambda pk: Membro.objects.get(pk=pk).nome)

    # Perfeição (4º-14º): colunas 4/9/14 — nada do 15 adiante; 4º some
    # porque não há candidato de Súplica neste mundo (coluna vazia sai)
    resposta_p = client.get('/iniciacao/', {'corpo': ['perfeicao']})
    colunas_p = {c['key']: c for c in resposta_p.context['colunas']}
    assert list(colunas_p) == ['pagamento', 'iniciacao']
    assert {'grau15', 'comunicacao'} & set(colunas_p) == set()

    # Kadosh (19/22/28/30): todos os marcos do corpo
    resposta_k = client.get('/iniciacao/', {'corpo': ['kadosh']})
    colunas_k = {c['key']: c for c in resposta_k.context['colunas']}
    assert list(colunas_k) == ['grau19', 'diploma', 'grau28', 'grau30']
    ids_k = [
        c['id']
        for col in colunas_k.values()
        for c in col['cards'] if c.get('tipo') == 'membro'
    ]
    assert set(ids_k) == {
        mundo['g19_ok'].pk, mundo['g22_ok'].pk,
        mundo['g28_ok'].pk, mundo['g30_ok'].pk,
    }

    # Múltiplos corpos: união das colunas
    resposta_m = client.get('/iniciacao/', {'corpo': ['capitulo', 'kadosh']})
    assert list(
        c['key'] for c in resposta_m.context['colunas']
    ) == ['grau15', 'comunicacao', 'grau19', 'diploma', 'grau28', 'grau30']

    # SC33: sem coluna de grau — aviso visível
    resposta_s = client.get('/iniciacao/', {'corpo': ['sc33']})
    assert 'Nenhuma coluna de grau pertence ao corpo selecionado.' in (
        resposta_s.content.decode()
    )


def test_card_membro_sem_vale_e_com_percentual_de_presenca():
    mundo = _mundo()
    # 1 presença, 2 faltas → 33% (< 50% → nome em vermelho)
    _presencas(mundo['g9_ok'], mundo['loja'], [True, False, False])
    client = _cliente(_admin())

    resposta = client.get('/iniciacao/')
    html = resposta.content.decode()
    colunas = {c['key']: c for c in resposta.context['colunas']}
    cards_g9 = [
        c for c in colunas['pagamento']['cards'] if c.get('tipo') == 'membro'
    ]

    assert resposta.status_code == 200
    # O card não exibe mais o nome do Vale
    assert all('vale' not in c for c in cards_g9)

    # Percentual de presença no card
    por_id = {c['id']: c for c in cards_g9}
    percentual = PresencaServico.percentuais([mundo['g9_ok']])[mundo['g9_ok'].pk]
    assert percentual == 33
    assert por_id[mundo['g9_ok'].pk]['presenca'] == 33
    # Sem chamada registrada conta como 0%
    assert por_id[mundo['g9_ok_inativo'].pk]['presenca'] == 0

    # Rótulo escrito: "Presença X%" nos dois casos
    assert 'Presença 33%' in html
    assert 'Presença 0%' in html

    # Nome em vermelho abaixo de 50%
    assert 'text-red-400' in html
    assert mundo['g9_ok'].nome in html


def test_relatorio_listagem_e_downloads():
    """Relatório = listagem dos irmãos filtrados (corpos/loja/vale), com
    download em PDF e Excel; contador substitui o 'Total de candidatos'."""
    mundo = _mundo()
    client = _cliente(_admin())
    membros_kadosh = {
        mundo['g19_ok'].nome,
        mundo['g22_ok'].nome,
        mundo['g28_ok'].nome,
        mundo['g30_ok'].nome,
    }

    # Listagem HTML do corpo Kadosh: irmãos dos marcos 19/22/28/30
    resposta = client.get('/iniciacao/relatorio/', {'corpo': ['kadosh']})
    conteudo = resposta.content.decode()
    assert resposta.status_code == 200
    for nome in membros_kadosh:
        assert nome in conteudo
    assert 'Kadosh' in conteudo
    # fora do corpo marcado não entra
    assert mundo['g9_ok'].nome not in conteudo

    # Excel: mesmas linhas, com capa de anexo
    resposta_x = client.get('/iniciacao/relatorio/xlsx/', {'corpo': ['kadosh']})
    assert resposta_x.status_code == 200
    assert resposta_x['Content-Type'] == (
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    assert resposta_x['Content-Disposition'] == (
        'attachment; filename="iniciacao-%s.xlsx"'
        % timezone.localdate().isoformat()
    )
    planilha = load_workbook(BytesIO(resposta_x.content)).active
    nomes_xlsx = [
        planilha.cell(row=n, column=2).value
        for n in range(2, planilha.max_row + 1)
    ]
    assert set(nomes_xlsx) == membros_kadosh

    # PDF: documento válido
    resposta_p = client.get('/iniciacao/relatorio/pdf/', {'corpo': ['kadosh']})
    assert resposta_p.status_code == 200
    assert resposta_p['Content-Type'] == 'application/pdf'
    assert resposta_p.content[:5] == b'%PDF-'

    # Loja de Perfeião: só as linhas que o quadro mostra (marcos 4º-14º)
    resposta_l = client.get('/iniciacao/relatorio/', {'loja': mundo['loja'].nome})
    conteudo_l = resposta_l.content.decode()
    assert mundo['g9_ok'].nome in conteudo_l
    assert mundo['g9_ok_inativo'].nome in conteudo_l
    assert 'Loja: ' + mundo['loja'].nome in conteudo_l

    # Contador de irmãos listados após os filtros (sem 'Total de candidatos')
    resposta_k = client.get('/iniciacao/')
    html_k = resposta_k.content.decode()
    assert 'Irmãos listados' in html_k
    assert 'Total de candidatos' not in html_k
