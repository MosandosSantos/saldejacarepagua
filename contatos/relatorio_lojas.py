"""In-memory PDF and XLSX renderers for filtered Loja reports."""
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

_COLUNAS = (
    ('Loja', 'nome'),
    ('Corpo', 'corpo'),
    ('Vale', 'prefeitura'),
    ('Presidente', 'deputado_mestre'),
    ('Celular do presidente', 'presidente_celular'),
    ('Membros', 'total_membros'),
    ('Situação', 'situacao_display'),
)

_PDF_NAVY = colors.HexColor('#0a1e3d')
_PDF_TEXT = colors.HexColor('#243447')
_PDF_MUTED = colors.HexColor('#526477')
_PDF_GRID = colors.HexColor('#cbd5e1')
_XLSX_NAVY = '0A1E3D'
_XLSX_GRID = 'CBD5E1'


def _texto(valor):
    return '' if valor is None else str(valor)


def _celula_texto(planilha, linha, coluna, valor):
    celula = planilha.cell(linha, coluna, value=_texto(valor))
    celula.data_type = 's'
    return celula


def gerar_xlsx_lojas(linhas, gerado_em):
    """Build a filtered Loja workbook in memory with safe text cells."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Lojas'
    workbook.properties.title = 'Relatório de Lojas'
    workbook.properties.subject = f'Relatório gerado em {gerado_em:%d/%m/%Y}'

    for coluna, (titulo, _campo) in enumerate(_COLUNAS, start=1):
        celula = _celula_texto(worksheet, 1, coluna, titulo)
        celula.font = Font(name='Manrope', size=10, bold=True, color='FFFFFF')
        celula.fill = PatternFill('solid', fgColor=_XLSX_NAVY)
        celula.alignment = Alignment(horizontal='center', vertical='center')

    for numero_linha, linha in enumerate(linhas, start=2):
        for coluna, (_titulo, campo) in enumerate(_COLUNAS, start=1):
            valor = linha.get(campo)
            if campo == 'total_membros':
                celula = worksheet.cell(numero_linha, coluna, value=valor or 0)
                celula.number_format = '0'
            else:
                celula = _celula_texto(worksheet, numero_linha, coluna, valor)
            celula.alignment = Alignment(
                horizontal='left' if coluna == 1 else 'center',
                vertical='center',
            )

    ultima_coluna = get_column_letter(len(_COLUNAS))
    worksheet.auto_filter.ref = f'A1:{ultima_coluna}{max(1, worksheet.max_row)}'
    worksheet.freeze_panes = 'A2'
    worksheet.sheet_view.showGridLines = True
    worksheet.sheet_view.zoomScale = 85
    worksheet.print_title_rows = '1:1'
    worksheet.page_setup.orientation = 'landscape'
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25

    for coluna, largura in enumerate((34, 18, 26, 28, 22, 12, 14), start=1):
        worksheet.column_dimensions[get_column_letter(coluna)].width = largura

    grade = Side(style='thin', color=_XLSX_GRID)
    for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row):
        for celula in row:
            celula.border = Border(
                left=grade, right=grade, top=grade, bottom=grade,
            )
            if celula.row > 1 and celula.row % 2 == 0:
                celula.fill = PatternFill('solid', fgColor='F3F6FA')

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _texto_pdf(valor):
    texto = _texto(valor).encode('latin-1', 'replace').decode('latin-1')
    return escape(texto)


def gerar_pdf_lojas(linhas, gerado_em):
    """Build a compact A4 landscape table that flows only vertically."""
    margem = 6 * mm
    titulo_style = ParagraphStyle(
        'LojasTitulo', fontName='Helvetica-Bold', fontSize=12,
        leading=14, textColor=_PDF_NAVY, spaceAfter=2 * mm,
    )
    subtitulo_style = ParagraphStyle(
        'LojasSubtitulo', fontName='Helvetica', fontSize=7,
        leading=9, textColor=_PDF_MUTED, spaceAfter=3 * mm,
    )
    header_style = ParagraphStyle(
        'LojasCabecalho', fontName='Helvetica-Bold', fontSize=7,
        leading=8, textColor=colors.white, alignment=1,
    )
    value_style = ParagraphStyle(
        'LojasValor', fontName='Helvetica', fontSize=7,
        leading=9, textColor=_PDF_TEXT,
    )
    empty_style = ParagraphStyle(
        'LojasVazio', fontName='Helvetica-Oblique', fontSize=8,
        leading=10, textColor=_PDF_MUTED, alignment=1,
    )

    cabecalho = [titulo for titulo, _campo in _COLUNAS]
    dados = [[Paragraph(_texto_pdf(titulo), header_style) for titulo in cabecalho]]
    for linha in linhas:
        dados.append([
            Paragraph(_texto_pdf(linha.get(campo)), value_style)
            for _titulo, campo in _COLUNAS
        ])

    larguras = [
        52 * mm, 28 * mm, 40 * mm, 44 * mm, 43 * mm, 22 * mm, 37 * mm,
    ]
    estilo = [
        ('BACKGROUND', (0, 0), (-1, 0), _PDF_NAVY),
        ('GRID', (0, 0), (-1, -1), 0.3, _PDF_GRID),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (5, 1), (-1, -1), 'CENTER'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]
    if not linhas:
        dados.append([
            Paragraph('Nenhuma loja encontrada para esses filtros.', empty_style),
        ] + [''] * (len(cabecalho) - 1))
        estilo.append(('SPAN', (0, 1), (-1, 1)))
    else:
        for numero_linha in range(2, len(dados), 2):
            estilo.append((
                'BACKGROUND', (0, numero_linha), (-1, numero_linha),
                colors.HexColor('#f3f6fa'),
            ))

    tabela = Table(
        dados,
        colWidths=larguras,
        repeatRows=1,
        splitByRow=1,
        hAlign='LEFT',
    )
    tabela.setStyle(TableStyle(estilo))

    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=margem,
        rightMargin=margem,
        topMargin=9 * mm,
        bottomMargin=9 * mm,
        title='Relatório de Lojas',
        author='GPRBRCRM',
    )
    documento.build([
        Paragraph('Relatório de Lojas', titulo_style),
        Paragraph(f'Gerado em {gerado_em:%d/%m/%Y}', subtitulo_style),
        tabela,
    ])
    return buffer.getvalue()
