"""Renderers in memory for the member-accompaniment report."""

from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

_PDF_NAVY = colors.HexColor('#0a1e3d')
_PDF_TEXT = colors.HexColor('#243447')
_PDF_MUTED = colors.HexColor('#526477')
_PDF_GRID = colors.HexColor('#cbd5e1')
_XLSX_NAVY = '0A1E3D'
_XLSX_ACCENT = '5B8DEF'
_XLSX_GRID = 'CBD5E1'


def _data_grau_mes_ano(valor):
    if valor is None:
        return '—'
    if isinstance(valor, datetime):
        valor = valor.date()
    return valor.strftime('%m/%Y')


def _texto_pdf(valor):
    texto = '' if valor is None else str(valor)
    texto = texto.encode('latin-1', 'replace').decode('latin-1')
    return escape(texto)


def _celula_texto(worksheet, row, column, value):
    cell = worksheet.cell(row=row, column=column, value='' if value is None else str(value))
    cell.data_type = 's'
    return cell


def gerar_xlsx_acompanhamento(linhas, graus_visiveis, gerado_em):
    """Return an XLSX workbook in memory with the same rows/columns as the table."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Acompanhamento'
    workbook.properties.title = 'Acompanhamento de Membros'
    workbook.properties.subject = f'Relatório gerado em {gerado_em:%d/%m/%Y}'

    cabecalhos = ['Membro', 'Presença'] + [
        f'{grau.numero}º' for grau in graus_visiveis
    ]
    for column, value in enumerate(cabecalhos, start=1):
        cell = _celula_texto(worksheet, 1, column, value)
        cell.font = Font(name='Manrope', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=_XLSX_NAVY)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_number, linha in enumerate(linhas, start=2):
        _celula_texto(worksheet, row_number, 1, linha['nome'])
        if linha['tem_presenca']:
            presenca = worksheet.cell(row=row_number, column=2, value=linha['presenca'])
            presenca.number_format = '0"%"'
        else:
            _celula_texto(worksheet, row_number, 2, '—')

        datas_por_grau = {
            item['numero']: item['data'] for item in linha['graus']
        }
        for offset, grau in enumerate(graus_visiveis, start=3):
            valor = datas_por_grau.get(grau.numero)
            if valor is None:
                _celula_texto(worksheet, row_number, offset, '—')
            else:
                cell = worksheet.cell(row=row_number, column=offset, value=valor)
                cell.number_format = 'mm/yyyy'
                cell.alignment = Alignment(horizontal='center')

    ultima_coluna = get_column_letter(len(cabecalhos))
    worksheet.auto_filter.ref = f'A1:{ultima_coluna}{max(1, worksheet.max_row)}'
    worksheet.freeze_panes = 'C2'
    worksheet.sheet_view.showGridLines = True
    worksheet.sheet_view.zoomScale = 75
    worksheet.print_title_rows = '1:1'
    worksheet.page_setup.orientation = 'landscape'
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.2
    worksheet.page_margins.right = 0.2

    worksheet.column_dimensions['A'].width = 38
    worksheet.column_dimensions['B'].width = 12
    for column in range(3, len(cabecalhos) + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 10

    thin_grid = Side(style='thin', color=_XLSX_GRID)
    for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row):
        for cell in row:
            cell.border = Border(
                left=thin_grid, right=thin_grid, top=thin_grid, bottom=thin_grid,
            )
            if cell.row > 1:
                cell.alignment = Alignment(
                    horizontal='left' if cell.column == 1 else 'center',
                    vertical='center',
                )
                if cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F3F6FA')
                if cell.column == 2 and isinstance(cell.value, int):
                    cell.font = Font(color=_XLSX_ACCENT, bold=True)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def gerar_pdf_acompanhamento(linhas, graus_visiveis, gerado_em):
    """Return a compact A4 landscape table with vertically flowing rows."""
    page_width, _page_height = landscape(A4)
    margem_horizontal = 5 * mm
    largura_util = page_width - 2 * margem_horizontal
    largura_membro = 38 * mm
    largura_presenca = 14 * mm
    quantidade_graus = len(graus_visiveis)
    largura_grau = (
        (largura_util - largura_membro - largura_presenca) / quantidade_graus
        if quantidade_graus
        else 0
    )
    fonte_grau = 4.8 if largura_grau < 8 * mm else 5.5

    titulo_style = ParagraphStyle(
        'AcompanhamentoTitulo', fontName='Helvetica-Bold', fontSize=11,
        leading=13, textColor=_PDF_NAVY, spaceAfter=2 * mm,
    )
    subtitulo_style = ParagraphStyle(
        'AcompanhamentoSubtitulo', fontName='Helvetica', fontSize=7,
        leading=9, textColor=_PDF_MUTED, spaceAfter=3 * mm,
    )
    header_style = ParagraphStyle(
        'AcompanhamentoCabecalho', fontName='Helvetica-Bold', fontSize=fonte_grau,
        leading=fonte_grau + 1, textColor=colors.white, alignment=1,
    )
    name_style = ParagraphStyle(
        'AcompanhamentoMembro', fontName='Helvetica', fontSize=6,
        leading=7, textColor=_PDF_TEXT,
    )
    value_style = ParagraphStyle(
        'AcompanhamentoValor', fontName='Helvetica', fontSize=fonte_grau,
        leading=fonte_grau + 1, textColor=_PDF_TEXT, alignment=1,
    )
    empty_style = ParagraphStyle(
        'AcompanhamentoVazio', fontName='Helvetica-Oblique', fontSize=7,
        leading=9, textColor=_PDF_MUTED, alignment=1,
    )

    cabecalho = ['Membro', 'Presença'] + [
        f'{grau.numero}º' for grau in graus_visiveis
    ]
    dados = [[Paragraph(_texto_pdf(valor), header_style) for valor in cabecalho]]
    for linha in linhas:
        presenca = f"{linha['presenca']}%" if linha['tem_presenca'] else '—'
        datas_por_grau = {
            item['numero']: item['data'] for item in linha['graus']
        }
        valores = [
            Paragraph(_texto_pdf(linha['nome']), name_style),
            Paragraph(_texto_pdf(presenca), value_style),
        ]
        valores.extend(
            Paragraph(_texto_pdf(_data_grau_mes_ano(datas_por_grau.get(grau.numero))), value_style)
            for grau in graus_visiveis
        )
        dados.append(valores)

    col_widths = [largura_membro, largura_presenca] + [
        largura_grau for _grau in graus_visiveis
    ]
    estilo = [
        ('BACKGROUND', (0, 0), (-1, 0), _PDF_NAVY),
        ('GRID', (0, 0), (-1, -1), 0.25, _PDF_GRID),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (1, 1), (-1, -1), 'CENTER'),
        ('LEFTPADDING', (0, 0), (-1, -1), 1),
        ('RIGHTPADDING', (0, 0), (-1, -1), 1),
        ('TOPPADDING', (0, 0), (-1, -1), 1.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 1.5),
    ]
    if not linhas:
        dados.append([
            Paragraph('Nenhum membro encontrado para esses filtros.', empty_style),
        ] + [''] * (len(cabecalho) - 1))
        estilo.append(('SPAN', (0, 1), (-1, 1)))
    else:
        for row_number in range(2, len(dados), 2):
            estilo.append((
                'BACKGROUND', (0, row_number), (-1, row_number),
                colors.HexColor('#f3f6fa'),
            ))

    tabela = Table(
        dados,
        colWidths=col_widths,
        repeatRows=1,
        splitByRow=1,
        hAlign='LEFT',
    )
    tabela.setStyle(TableStyle(estilo))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=margem_horizontal,
        rightMargin=margem_horizontal,
        topMargin=9 * mm,
        bottomMargin=10 * mm,
        title='Acompanhamento de Membros',
        author='GPRBRCRM',
    )
    story = [
        Paragraph('Acompanhamento de Membros', titulo_style),
        Paragraph(f'Gerado em {gerado_em:%d/%m/%Y}', subtitulo_style),
        tabela,
    ]
    doc.build(story)
    return buffer.getvalue()
